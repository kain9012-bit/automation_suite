"""파일·엑셀 도구 회귀 확인(평범한 스크립트).

    python tests/regression/test_file_tools.py

시험용 입력(엑셀·CSV·ZIP·폴더·작은 웹 서버)은 실행할 때 임시 폴더에 만든다.
도구는 앱과 같은 방식으로 bridge/runner.py 에 --tool/--payload-b64 로 넘겨 돌린다.
윈도에서만 생기는 경우(드라이브 루트, 정션)는 도우미 함수를 직접 부르거나
os.walk/os.path 를 윈도처럼 바꿔 끼워(monkeypatch) 확인한다.
하나라도 틀리면 종료 코드 1.
"""
from __future__ import annotations

import base64
import contextlib
import http.server
import importlib
import json
import ntpath
import os
import socket
import struct
import subprocess
import sys
import tempfile
import threading
import traceback
import zipfile
import zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TOOLS = ROOT / "tools"
BRIDGE = ROOT / "bridge"
PY = sys.executable
PROGRESS = "@@JBEDU_PROGRESS@@"

ENV = dict(os.environ)
ENV["JBEDU_TOOLS_ROOT"] = str(TOOLS)
ENV["PYTHONIOENCODING"] = "utf-8"
ENV["NO_PROXY"] = ENV["no_proxy"] = "127.0.0.1,localhost"
os.environ["JBEDU_TOOLS_ROOT"] = str(TOOLS)
os.environ["NO_PROXY"] = os.environ["no_proxy"] = "127.0.0.1,localhost"

RESULTS: list[tuple[bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> bool:
    RESULTS.append((bool(ok), name))
    mark = "✔" if ok else "✗"
    print(f"{mark} {name}" + ("" if ok or not detail else f"\n    → {detail}"), flush=True)
    return bool(ok)


def attempt(func) -> bool:
    """도우미 함수 확인. 수정 전 코드처럼 함수가 없거나 인자가 다르면 False."""
    try:
        return bool(func())
    except Exception as exc:  # noqa: BLE001
        print(f"    (예외: {type(exc).__name__}: {exc})")
        return False


def case(func):
    """한 묶음의 확인. 예외가 나면(수정 전 코드 등) 그 묶음을 실패로 적고 다음으로 넘어간다."""
    def wrapper():
        try:
            func()
        except Exception as exc:  # noqa: BLE001
            check(f"{func.__name__}: 예외 없이 끝남", False, f"{type(exc).__name__}: {exc}\n{traceback.format_exc(limit=3)}")
    wrapper.__name__ = func.__name__
    return wrapper


class Run:
    def __init__(self, code: int, result: dict | None, stderr: str, stdout_lines: list[str]):
        self.code = code
        self.result = result or {}
        self.stderr = stderr
        self.stdout_lines = stdout_lines

    @property
    def progress(self) -> list[str]:
        return [line for line in self.stdout_lines if line.startswith(PROGRESS)]


def run_tool(tool: str, payload: dict, timeout: int = 180) -> Run:
    b64 = base64.b64encode(json.dumps(payload, ensure_ascii=False).encode("utf-8")).decode("ascii")
    proc = subprocess.run(
        [PY, str(BRIDGE / "runner.py"), "--tool", tool, "--payload-b64", b64],
        capture_output=True, env=ENV, timeout=timeout,
    )
    stdout = proc.stdout.decode("utf-8", "replace")
    stderr = proc.stderr.decode("utf-8", "replace")
    lines = [line for line in stdout.splitlines() if line.strip()]
    others = [line for line in lines if not line.startswith(PROGRESS)]
    result = None
    if proc.returncode == 0:
        # 결과 계약: 진행 줄 말고는 마지막 JSON 한 줄뿐이어야 한다.
        if len(others) != 1:
            check(f"{tool}: stdout 은 진행 줄 + JSON 한 줄", False, repr(others[:5]))
        try:
            result = json.loads(others[-1])
        except Exception:
            result = None
    return Run(proc.returncode, result, stderr, lines)


def in_process(module: str, tool_dir: str | None = None):
    if tool_dir:
        path = str(TOOLS / tool_dir)
        if path not in sys.path:
            sys.path.insert(0, path)
    if str(BRIDGE) not in sys.path:
        sys.path.insert(0, str(BRIDGE))
    return importlib.import_module(module)


@contextlib.contextmanager
def windows_junctions(*junctions: Path):
    """리눅스 심볼릭 링크를 윈도 정션처럼 보이게 한다.

    윈도에서 정션은 os.path.islink 가 False, os.path.isjunction 이 True 이고,
    os.walk 가 그 안으로 따라 들어가며 Path.resolve() 는 실제 위치(밖)로 풀린다.
    """
    names = {os.path.normcase(os.path.abspath(p)) for p in junctions}
    orig_walk, orig_islink = os.walk, os.path.islink
    orig_isjunction = getattr(os.path, "isjunction", None)

    def walk(top, topdown=True, onerror=None, followlinks=False):
        return orig_walk(top, topdown, onerror, True)

    def islink(path):
        if os.path.normcase(os.path.abspath(os.fspath(path))) in names:
            return False
        return orig_islink(path)

    def isjunction(path):
        return os.path.normcase(os.path.abspath(os.fspath(path))) in names

    os.walk, os.path.islink, os.path.isjunction = walk, islink, isjunction
    try:
        yield
    finally:
        os.walk, os.path.islink = orig_walk, orig_islink
        if orig_isjunction is None:
            del os.path.isjunction
        else:
            os.path.isjunction = orig_isjunction


# ---------------------------------------------------------------- 엑셀·ZIP 만들기 도우미

def add_cached_values(xlsx: Path, sheet_index: int, values: dict[str, object]) -> None:
    """openpyxl 이 쓴 xlsx 는 함수의 계산 결과(<v>)가 없다. 엑셀이 저장한 것처럼 넣어 준다."""
    import re

    member = f"xl/worksheets/sheet{sheet_index}.xml"
    with zipfile.ZipFile(xlsx) as zin:
        items = {name: zin.read(name) for name in zin.namelist()}
    xml = items[member].decode("utf-8")
    for coord, value in values.items():
        pattern = re.compile(rf'(<c r="{coord}"[^>]*>)(<f>.*?</f>)(<v\s*/>|<v></v>)?')
        if isinstance(value, str):
            repl = lambda m, v=value: m.group(1).replace("<c ", '<c t="str" ', 1) + m.group(2) + f"<v>{v}</v>"
        else:
            repl = lambda m, v=value: m.group(1) + m.group(2) + f"<v>{v}</v>"
        xml, count = pattern.subn(repl, xml)
        assert count == 1, (coord, xml[:500])
    items[member] = xml.encode("utf-8")
    with zipfile.ZipFile(xlsx, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items.items():
            zout.writestr(name, data)


_CRC_TABLE = []
for _i in range(256):
    _c = _i
    for _ in range(8):
        _c = (_c >> 1) ^ 0xEDB88320 if _c & 1 else _c >> 1
    _CRC_TABLE.append(_c)


def _zipcrypto_encrypt(data: bytes, pwd: bytes, crc: int) -> bytes:
    keys = [0x12345678, 0x23456789, 0x34567890]

    def crc32(value, byte):
        return _CRC_TABLE[(value ^ byte) & 0xFF] ^ (value >> 8)

    def update(byte):
        keys[0] = crc32(keys[0], byte)
        keys[1] = (keys[1] + (keys[0] & 0xFF)) & 0xFFFFFFFF
        keys[1] = (keys[1] * 134775813 + 1) & 0xFFFFFFFF
        keys[2] = crc32(keys[2], keys[1] >> 24)

    def stream_byte():
        temp = (keys[2] | 2) & 0xFFFF
        return ((temp * (temp ^ 1)) >> 8) & 0xFF

    for byte in pwd:
        update(byte)
    header = bytes([0x5A, 0x11, 0x42, 0x07, 0x99, 0x31, 0x63, 0x0D, 0x2E, 0x44, 0x71]) + bytes([(crc >> 24) & 0xFF])
    out = bytearray()
    for byte in header + data:
        out.append(byte ^ stream_byte())
        update(byte)
    return bytes(out)


def raw_zip(path: Path, entries: list[dict]) -> None:
    """zipfile 이 만들지 못하는 ZIP(옛 ZIP 암호, AES 표시, UTF-8 표시 없는 UTF-8 이름)을 직접 쓴다.

    entry: {"name": bytes, "data": bytes, "flags": int, "method": int, "pwd": bytes | None}
    """
    body = bytearray()
    central = bytearray()
    for entry in entries:
        name, data = entry["name"], entry["data"]
        flags, method = entry.get("flags", 0), entry.get("method", 0)
        crc = zlib.crc32(data) & 0xFFFFFFFF
        stored = data
        if entry.get("pwd") is not None:
            flags |= 0x1
            stored = _zipcrypto_encrypt(data, entry["pwd"], crc)
        offset = len(body)
        body += struct.pack("<4sHHHHHIIIHH", b"PK\x03\x04", 20, flags, method, 0, 0x5A21, crc,
                            len(stored), len(data), len(name), 0) + name + stored
        central += struct.pack("<4sHHHHHHIIIHHHHHII", b"PK\x01\x02", 20, 20, flags, method, 0, 0x5A21, crc,
                               len(stored), len(data), len(name), 0, 0, 0, 0, 0, offset) + name
    eocd = struct.pack("<4sHHHHIIH", b"PK\x05\x06", 0, 0, len(entries), len(entries), len(central), len(body), 0)
    path.write_bytes(bytes(body + central + eocd))


def load(path: Path, data_only: bool = False):
    from openpyxl import load_workbook
    return load_workbook(path, data_only=data_only)


# ---------------------------------------------------------------- 1. folder_unpacker

@case
def test_folder_unpacker_drive_root():
    svc = in_process("folder_unpacker_service", "folder_unpacker")
    roots = ["D:\\", "d:/", "\\\\server\\share\\", "\\\\server\\share"]
    not_roots = ["D:\\자료", "\\\\server\\share\\학교"]
    check("folder_unpacker: D:\\ · \\\\서버\\공유 를 드라이브 루트로 알아봄", attempt(
        lambda: all(svc.is_drive_root(p, pathmod=ntpath) for p in roots)
        and not any(svc.is_drive_root(p, pathmod=ntpath) for p in not_roots)
    ))

    # 드라이브 루트로 판정되면 아무 파일도 옮기지 않고 한국어 오류로 멈춰야 한다.
    with tempfile.TemporaryDirectory() as tmp:
        pick = Path(tmp) / "pick"
        (pick / "sub").mkdir(parents=True)
        (pick / "sub" / "a.txt").write_text("a", encoding="utf-8")
        original = svc.is_drive_root
        svc.is_drive_root = lambda path, pathmod=os.path: True
        error = ""
        try:
            svc.unpack_folders([pick], recursive=True)
        except Exception as exc:  # noqa: BLE001
            error = str(exc)
        finally:
            svc.is_drive_root = original
        check(
            "folder_unpacker: 드라이브 루트를 고르면 옮기지 않고 거절",
            "드라이브 루트" in error and (pick / "sub" / "a.txt").exists(),
            f"error={error!r}, 파일 남음={(pick / 'sub' / 'a.txt').exists()}",
        )


@case
def test_folder_unpacker_junction():
    runner = in_process("runner")
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        outside = base / "outside"
        outside.mkdir()
        (outside / "secret.txt").write_text("밖에 있는 파일", encoding="utf-8")
        (outside / "빈폴더").mkdir()
        pick = base / "parent" / "pick"
        (pick / "sub").mkdir(parents=True)
        (pick / "sub" / "a.txt").write_text("a", encoding="utf-8")
        junction = pick / "junc"
        os.symlink(outside, junction, target_is_directory=True)
        with windows_junctions(junction):
            runner.run_folder_unpacker({"inputs": [str(pick)], "confirmed": True, "recursive": True})
        check(
            "folder_unpacker: 정션 안(고른 폴더 밖) 파일은 옮기지 않음",
            (outside / "secret.txt").exists() and (outside / "빈폴더").is_dir()
            and not (base / "parent" / "secret.txt").exists(),
            f"outside={sorted(p.name for p in outside.iterdir())}",
        )
        check("folder_unpacker: 안쪽 파일은 그대로 꺼냄", (base / "parent" / "a.txt").exists())


# ---------------------------------------------------------------- 2. empty_folder_cleaner

@case
def test_empty_folder_cleaner_junction():
    extra = in_process("extra_tools")
    sys.path.insert(0, str(TOOLS / "empty_folder_cleaner"))
    for name in list(sys.modules):
        if name == "services" or name.startswith("services."):
            del sys.modules[name]
    svc = importlib.import_module("services.empty_folder_cleaner_service")
    check("empty_folder_cleaner: 드라이브·공유 루트 판정", attempt(
        lambda: svc.is_drive_root("\\\\server\\share", pathmod=ntpath) and svc.is_drive_root("E:\\", pathmod=ntpath)
        and not svc.is_drive_root("E:\\자료", pathmod=ntpath)
    ))

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        outside = base / "outside"
        (outside / "e").mkdir(parents=True)
        root = base / "root"
        (root / "a" / "b").mkdir(parents=True)
        (root / "keep").mkdir()
        (root / "keep" / "x.txt").write_text("x", encoding="utf-8")
        junction = root / "junc"
        os.symlink(outside, junction, target_is_directory=True)
        with windows_junctions(junction):
            result = extra.run_empty_folder_cleaner({"inputs": [str(root)], "confirmed": True})
        check(
            "empty_folder_cleaner: 정션 너머(기준 폴더 밖) 빈 폴더는 지우지 않음",
            (outside / "e").is_dir() and outside.is_dir(),
            f"outside/e 남음={(outside / 'e').is_dir()}",
        )
        check("empty_folder_cleaner: 정션 자체도 지우지 않음", os.path.lexists(junction))
        check(
            "empty_folder_cleaner: 기준 폴더 안 빈 폴더는 지움",
            not (root / "a").exists() and (root / "keep" / "x.txt").exists(),
            json.dumps(result.get("result", {}), ensure_ascii=False)[:300],
        )


# ---------------------------------------------------------------- 3·4. excel_split

def _grade_book(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "성적"
    ws.append(["이름", "반", "국어", "수학", "합계", "가중", "두배", "메모"])
    data = [("홍길동", "1반", 90, 80), ("김철수", "2반", 70, 60), ("이영희", "1반", 85, 95), ("박민수", "2반", 50, 40)]
    for idx, (name, klass, kor, math) in enumerate(data, start=2):
        ws.append([name, klass, kor, math, f"=C{idx}+D{idx}", f"=기준!A1*C{idx}", f"=C{idx}*2", "=== 확인 필요"])
        ws[f"H{idx}"].data_type = "s"  # 엑셀에서 글자로 적은 '=== 확인 필요'
    ref = wb.create_sheet("기준")
    ref["A1"] = 2
    wb.save(path)
    cached = {}
    for idx, (_, _, kor, math) in enumerate(data, start=2):
        cached[f"E{idx}"] = kor + math
        cached[f"F{idx}"] = 2 * kor
        # G 열은 결과값을 일부러 비워 둔다(엑셀이 아닌 프로그램이 저장한 파일처럼).
    add_cached_values(path, 1, cached)


@case
def test_excel_split_formulas():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        src = base / "성적.xlsx"
        _grade_book(src)
        out = base / "out"
        run = run_tool("excel_split", {
            "inputs": [str(src)], "output": str(out), "mode": "column", "sheet_name": "성적",
            "header_row": 1, "split_column": 2, "rows_per_file": 1000, "skip_empty_key": True,
        })
        check("excel_split: 열 기준 분할 성공", run.code == 0, run.stderr[-300:])
        target = out / "성적_2반.xlsx"
        ws = load(target)["성적"]
        # 2반 = 원본 3행(김철수)·5행(박민수) → 새 파일 2·3행
        check("excel_split: 합계 함수는 계산된 값(김철수 130)", ws["E2"].value == 130, repr(ws["E2"].value))
        check("excel_split: 합계 함수는 계산된 값(박민수 90)", ws["E3"].value == 90, repr(ws["E3"].value))
        check("excel_split: 다른 시트를 가리키던 함수도 값(140)", ws["F2"].value == 140, repr(ws["F2"].value))
        check("excel_split: 결과값 없는 함수는 새 위치로 옮김(=C2*2)", ws["G2"].value == "=C2*2", repr(ws["G2"].value))
        check("excel_split: '='로 시작하는 글자는 글자로", ws["H2"].data_type == "s" and ws["H2"].value == "=== 확인 필요",
              f"{ws['H2'].data_type} {ws['H2'].value!r}")


@case
def test_excel_split_output_folder():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        src_dir = base / "src"
        src_dir.mkdir()
        src = src_dir / "2026.03 학생명단.xlsx"
        _grade_book(src)
        dotted = base / "2026.03 자료"
        run = run_tool("excel_split", {"inputs": [str(src)], "output": str(dotted), "mode": "sheet"})
        made = sorted(p.name for p in dotted.glob("*.xlsx")) if dotted.is_dir() else []
        check("excel_split: 점이 든 저장 폴더('2026.03 자료') 안에 저장", run.code == 0 and len(made) == 2,
              f"code={run.code} made={made} base={sorted(p.name for p in base.iterdir())}")

        run = run_tool("excel_split", {"inputs": [str(src)], "output": "", "mode": "sheet"})
        default = src_dir / "2026.03 학생명단_분할"
        made = sorted(p.name for p in default.glob("*.xlsx")) if default.is_dir() else []
        check("excel_split: 저장 폴더를 비우면 '(파일명)_분할' 폴더에 저장", run.code == 0 and len(made) == 2,
              f"code={run.code} src_dir={sorted(p.name for p in src_dir.iterdir())}")


# ---------------------------------------------------------------- 5. zip_batch_extractor

@case
def test_zip_passwords_and_errors():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        locked = base / "locked.zip"
        raw_zip(locked, [{"name": b"secret.txt", "data": b"top secret", "pwd": b"abcd"}])
        out = base / "out"
        out.mkdir()
        run = run_tool("zip_batch_extractor", {"inputs": [str(locked)], "output": str(out), "password": "wrong"})
        check("zip: 틀린 암호 → '암호가 맞지 않습니다' 오류", run.code != 0 and "암호가 맞지 않" in run.stderr,
              f"code={run.code} msg={run.result.get('message')!r} err={run.stderr[-200:]!r}")
        check("zip: 틀린 암호로 빈 폴더를 남기지 않음", not any(out.iterdir()), str(list(out.iterdir())))

        out = base / "out2"
        out.mkdir()
        run = run_tool("zip_batch_extractor", {"inputs": [str(locked)], "output": str(out), "password": "abcd"})
        secret = out / "locked" / "secret.txt"
        check("zip: 맞는 암호로 풀림", run.code == 0 and secret.exists() and secret.read_bytes() == b"top secret",
              run.stderr[-200:])

        korean = base / "한글암호.zip"
        raw_zip(korean, [{"name": b"k.txt", "data": b"korean pwd", "pwd": "비밀번호".encode("cp949")}])
        run = run_tool("zip_batch_extractor", {"inputs": [str(korean)], "output": str(out), "password": "비밀번호"})
        check("zip: cp949 로 걸린 한글 암호도 풀림", run.code == 0 and (out / "한글암호" / "k.txt").exists(),
              f"code={run.code} err={run.stderr[-200:]!r}")


@case
def test_zip_corrupt_and_unsupported():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        good = base / "good.zip"
        with zipfile.ZipFile(good, "w") as zf:
            zf.writestr("ok.txt", "ok")
        broken = base / "broken.zip"
        broken.write_bytes(b"this is not a zip file at all" * 10)
        aes = base / "aes.zip"
        raw_zip(aes, [{"name": b"a.txt", "data": b"\x00" * 40, "method": 99, "flags": 0x1}])
        mac = base / "mac.zip"
        raw_zip(mac, [{"name": "한글문서.txt".encode("utf-8"), "data": b"mac"}])
        out = base / "out"
        out.mkdir()
        run = run_tool("zip_batch_extractor", {"inputs": [str(good), str(broken), str(aes), str(mac)], "output": str(out)})
        failed = {item.get("name"): item.get("error", "") for item in run.result.get("failed", [])}
        check("zip: 일부 실패해도 나머지는 풂", run.code == 0 and (out / "good" / "ok.txt").exists(), run.stderr[-200:])
        check("zip: 손상된 ZIP → failed 에 '손상된 ZIP 파일입니다'", "손상된 ZIP" in failed.get("broken.zip", ""), repr(failed))
        check("zip: AES 등 → failed 에 '지원하지 않는 압축 방식'", "지원하지 않는 압축 방식" in failed.get("aes.zip", ""), repr(failed))
        check("zip: 실패가 있으면 완전 성공처럼 말하지 않음", "풀지 못했" in str(run.result.get("message", "")),
              repr(run.result.get("message")))
        check("zip: 실패한 ZIP 의 빈 폴더를 남기지 않음", not (out / "aes").exists() and not (out / "broken").exists(),
              str(sorted(p.name for p in out.iterdir())))
        check("zip: 맥 ZIP(UTF-8 표시 없음) 한글 이름", (out / "mac" / "한글문서.txt").exists(),
              str(sorted(p.name for p in (out / "mac").iterdir())) if (out / "mac").exists() else "no folder")
        check("zip: 진행 상황을 내보냄", any("good.zip" in line for line in run.progress), repr(run.progress[:3]))


# ---------------------------------------------------------------- 6. homepage_post_collector

LIST_HTML = """<html><head><meta charset="utf-8"></head><body>
<table><thead><tr><th>번호</th><th>제목</th><th>작성자</th></tr></thead><tbody>
<tr><td>1</td><td><a href="/board/view.jbe?dataSid=1">공지\x08테스트</a></td><td>교무\x0b부</td></tr>
<tr><td>2</td><td><a href="/board/view.jbe?dataSid=2">=== 안내</a></td><td>행정실</td></tr>
</tbody></table></body></html>"""
DETAIL_HTML = """<html><head><meta charset="utf-8"></head><body><h3>{title}</h3>
<div class="content">본문\x0c내용</div>
<a href="/download.do?id={sid}">다운로드</a> <a href="/filedown.do?id={sid}x">미리보기</a>
</body></html>"""


class _BoardHandler(http.server.BaseHTTPRequestHandler):
    def log_message(self, *args):  # 조용히
        pass

    def _send(self, body: bytes, ctype: str, extra: dict | None = None):
        self.send_response(200)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        for key, value in (extra or {}).items():
            self.send_header(key, value)
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        if self.path.startswith("/board/list"):
            self._send(LIST_HTML.encode("utf-8"), "text/html; charset=utf-8")
        elif self.path.startswith("/board/view"):
            sid = self.path.rsplit("=", 1)[-1]
            self._send(DETAIL_HTML.format(title="글", sid=sid).encode("utf-8"), "text/html; charset=utf-8")
        elif self.path.startswith("/download"):
            self._send(b"HWP", "application/octet-stream",
                       {"Content-Disposition": "attachment; filename*=UTF-8''%EA%B3%84%ED%9A%8D%EC%84%9C.hwp"})
        elif self.path.startswith("/filedown"):
            self._send(b"%PDF", "application/octet-stream", {"Content-Disposition": 'attachment; filename="report.pdf"'})
        else:
            self.send_error(404)


@case
def test_homepage_collector():
    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _BoardHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            run = run_tool("homepage_post_collector", {
                "board_url": f"http://127.0.0.1:{server.server_port}/board/list.jbe",
                "output": str(out), "start_page": 1, "end_page": 1, "max_posts": 0,
                "collect_body": True, "download_files": True, "create_excel": True,
            })
            check("homepage: 제목에 제어 문자가 있어도 엑셀 저장 성공", run.code == 0, run.stderr[-300:])
            excel = next(out.glob("게시글_취합_*.xlsx"), None)
            if excel is not None:
                ws = load(excel)["게시글"]
                titles = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
                check("homepage: 제어 문자만 빠지고 제목은 남음", "공지테스트" in titles, repr(titles))
                cell = next(ws.cell(row=r, column=2) for r in range(2, ws.max_row + 1)
                            if str(ws.cell(row=r, column=2).value).startswith("==="))
                check("homepage: '='로 시작하는 제목은 글자로", cell.data_type == "s", cell.data_type)
            else:
                check("homepage: 엑셀 파일이 만들어짐", False, str(list(out.iterdir())))
            saved = sorted(p.name for p in out.rglob("*") if p.is_file() and p.suffix != ".xlsx")
            check("homepage: '다운로드' 링크 첨부는 서버 파일 이름(계획서.hwp)으로 저장",
                  "계획서.hwp" in saved and "다운로드" not in saved, repr(saved))
            check("homepage: '미리보기' 링크 첨부는 확장자 있는 이름(report.pdf)으로 저장",
                  "report.pdf" in saved and "미리보기" not in saved, repr(saved))
    finally:
        server.shutdown()

    sock = socket.socket()
    sock.bind(("127.0.0.1", 0))
    port = sock.getsockname()[1]
    sock.close()
    with tempfile.TemporaryDirectory() as tmp:
        run = run_tool("homepage_post_collector", {
            "board_url": f"http://127.0.0.1:{port}/board/list.jbe", "output": tmp,
            "start_page": 1, "end_page": 2, "create_excel": True,
        })
        check("homepage: 모든 페이지 요청이 실패하면 이유와 함께 오류",
              run.code != 0 and "열지 못" in run.stderr and "접속" in run.stderr,
              f"code={run.code} msg={run.result.get('message')!r} err={run.stderr[-200:]!r}")


# ---------------------------------------------------------------- 7. excel_merge

def _merge_inputs(base: Path) -> list[Path]:
    from openpyxl import Workbook

    paths = []
    for idx in (1, 2):
        wb = Workbook()
        ws = wb.active
        ws.title = "명단"
        ws.append(["지역", "메모", "번호"])
        ws.append([f"전주시\n완산구{idx}", "  들여쓴 글", idx])
        ws.append(["군산시", "=== 확인 필요", idx + 10])
        ws["B3"].data_type = "s"  # 엑셀에서 글자로 적은 '=== 확인 필요'
        path = base / f"학교{idx}.xlsx"
        wb.save(path)
        paths.append(path)
    csv_path = base / "카드.csv"
    csv_path.write_text("지역,메모,번호\n익산시,카드,1234567890123456789\n김제시,짧은,12345\n", encoding="utf-8-sig")
    paths.append(csv_path)
    return paths


@case
def test_excel_merge_values_unchanged():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        inputs = _merge_inputs(base)
        out = base / "병합.xlsx"
        run = run_tool("excel_merge", {"inputs": [str(p) for p in inputs], "output": str(out),
                                       "mode": "workbook_sheets_to_one"})
        check("excel_merge: 한 시트로 병합 성공", run.code == 0, run.stderr[-300:])
        ws = load(out)["병합"]
        values = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        check("excel_merge: 칸 안 줄바꿈 유지", values[0][0] == "전주시\n완산구1", repr(values[0][0]))
        check("excel_merge: 앞 공백 유지", values[0][1] == "  들여쓴 글", repr(values[0][1]))
        memo = ws.cell(row=3, column=2)
        check("excel_merge: '=== 확인 필요'는 글자로", memo.data_type == "s" and memo.value == "=== 확인 필요",
              f"{memo.data_type} {memo.value!r}")
        card = next(row[2] for row in values if row[0] == "익산시")
        short = next(row[2] for row in values if row[0] == "김제시")
        check("excel_merge: CSV 의 16자리 넘는 숫자는 글자 그대로", card == "1234567890123456789", repr(card))
        check("excel_merge: CSV 의 보통 숫자는 숫자로", short == 12345, repr(short))
        popup = str((run.result.get("popup") or {}).get("message", ""))
        check("excel_merge: '='로 시작하는 글자를 함수로 세지 않음", "함수" not in popup, popup[:200])

        out2 = base / "시트별.xlsx"
        run = run_tool("excel_merge", {"inputs": [str(p) for p in inputs[:2]], "output": str(out2),
                                       "mode": "files_to_sheets"})
        cell = load(out2).worksheets[0]["B3"]
        check("excel_merge: 시트 그대로 옮기기에서도 '=' 글자는 글자로",
              run.code == 0 and cell.data_type == "s" and cell.value == "=== 확인 필요",
              f"{cell.data_type} {cell.value!r} {run.stderr[-200:]}")


# ---------------------------------------------------------------- 8. file_inventory

@case
def test_file_inventory():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        folder = base / "대상"
        (folder / "잠김").mkdir(parents=True)
        (folder / "=1+1.txt").write_text("x", encoding="utf-8")
        out = base / "현황.xlsx"
        run = run_tool("file_inventory", {"inputs": [str(folder)], "output": str(out), "recursive": True})
        check("file_inventory: 현황표 저장", run.code == 0, run.stderr[-300:])
        ws = load(out)["파일_폴더_목록"]
        cell = next(ws.cell(row=r, column=3) for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=3).value == "=1+1.txt")
        check("file_inventory: '='로 시작하는 파일 이름은 글자로", cell.data_type == "s", cell.data_type)

        # 읽을 수 없는 폴더(권한 없음)는 오류 기록에 남아야 한다. 관리자(root)로 돌면 권한이 안 막히므로
        # os.scandir 를 바꿔 끼워 그 폴더에서 PermissionError 를 낸다.
        runner = in_process("runner")
        locked = str(folder / "잠김")
        original = os.scandir

        def scandir(path="."):
            if os.path.abspath(os.fspath(path)) == locked:
                raise PermissionError(13, "액세스가 거부되었습니다", os.fspath(path))
            return original(path)

        os.scandir = scandir
        try:
            result = runner.run_file_inventory({"inputs": [str(folder)], "output": str(base / "현황2.xlsx"), "recursive": True})
        finally:
            os.scandir = original
        errors = [[c.value for c in row] for row in load(base / "현황2.xlsx")["오류_기록"].iter_rows(min_row=2)]
        check("file_inventory: 읽지 못한 폴더를 오류 기록에 남김",
              result.get("error_count") == 1 and any(locked == row[0] for row in errors), repr(errors))
        check("file_inventory: 오류가 있으면 메시지에 알림", "읽지 못한" in str(result.get("message")), repr(result.get("message")))


# ---------------------------------------------------------------- 9. rename_files

@case
def test_rename_files():
    svc = in_process("rename_files_service", "rename_files")
    bad = ["a<b.txt", "a:b.txt", 'a"b.txt', "a|b.txt", "a?b.txt", "a*b.txt", "CON.txt", "con", "Com1.pdf",
           "LPT9", "abc.", "abc ", "", "a/b.txt", "a\\b.txt"]
    good = ["보고서.txt", "CONTENT.txt", "a.b.c.txt"]
    check("rename_files: 윈도에서 못 쓰는 이름을 한국어 이유로 걸러냄", attempt(
        lambda: all(svc.invalid_name_reason(n) for n in bad) and not any(svc.invalid_name_reason(n) for n in good)
    ))

    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        files = [base / "ax.txt", base / "x.txt", base / "CONx.txt"]
        for path in files:
            path.write_text("t", encoding="utf-8")
        run = run_tool("rename_files", {"inputs": [str(p) for p in files], "confirmed": True,
                                        "mode": "replace", "before": "x", "after": ""})
        failed = {item["name"]: item["error"] for item in run.result.get("failed", [])}
        check("rename_files: 쓸 수 없는 이름은 건너뛰고 나머지는 바꿈",
              run.code == 0 and (base / "a.txt").exists() and (base / "x.txt").exists() and (base / "CONx.txt").exists(),
              f"code={run.code} files={sorted(p.name for p in base.iterdir())} err={run.stderr[-200:]!r}")
        check("rename_files: 실패한 파일을 failed 에 한국어 이유로", set(failed) == {"x.txt", "CONx.txt"}
              and all(any("가" <= ch <= "힣" for ch in err) for err in failed.values()), repr(failed))
        check("rename_files: 실패가 있으면 메시지에 알림", "바꾸지 못했" in str(run.result.get("message")),
              repr(run.result.get("message")))

        files = [base / "a.txt", base / "x.txt"]
        run = run_tool("rename_files", {"inputs": [str(p) for p in files], "confirmed": True,
                                        "mode": "suffix", "after": "/백업"})
        check("rename_files: '/'가 든 새 이름에서 전체가 멈추지 않음",
              run.code == 0 and len(run.result.get("failed", [])) == 2 and (base / "a.txt").exists(),
              f"code={run.code} err={run.stderr[-200:]!r}")


# ---------------------------------------------------------------- 10. 열 수 없는 엑셀

@case
def test_unreadable_workbook():
    with tempfile.TemporaryDirectory() as tmp:
        base = Path(tmp)
        drm = base / "보호문서.xlsx"
        # 암호를 건 xlsx 와 DRM 문서는 ZIP 이 아니라 OLE(복합 문서) 머리를 가진다.
        drm.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 1024)
        want = "엑셀 파일을 열 수 없습니다(암호·DRM 보호 또는 손상): 보호문서.xlsx"
        run = run_tool("excel_split__analyze", {"inputs": [str(drm)]})
        check("excel_split: 열 수 없는 파일 → 한국어 이유와 파일 이름", run.code != 0 and want in run.stderr,
              run.stderr[-200:])
        run = run_tool("excel_merge", {"inputs": [str(drm)], "output": str(base / "r.xlsx"),
                                       "mode": "workbook_sheets_to_one"})
        check("excel_merge: 열 수 없는 파일만 있으면 그 이유를 알림", run.code != 0 and want in run.stderr,
              run.stderr[-200:])
        run = run_tool("excel_merge", {"inputs": [str(drm)], "output": str(base / "r2.xlsx"), "mode": "files_to_sheets"})
        check("excel_merge(시트 그대로): 열 수 없는 이유를 알림", run.code != 0 and want in run.stderr, run.stderr[-200:])
        good = _merge_inputs(base)[0]
        run = run_tool("excel_merge", {"inputs": [str(good), str(drm)], "output": str(base / "r3.xlsx"),
                                       "mode": "workbook_sheets_to_one"})
        failed = run.result.get("failed", [])
        check("excel_merge: 일부만 못 열면 failed 에 담고 계속",
              run.code == 0 and any(item.get("name") == "보호문서.xlsx" for item in failed)
              and "읽지 못해" in str(run.result.get("message")), f"code={run.code} failed={failed!r}")


def main() -> int:
    tests = [
        test_folder_unpacker_drive_root,
        test_folder_unpacker_junction,
        test_empty_folder_cleaner_junction,
        test_excel_split_formulas,
        test_excel_split_output_folder,
        test_zip_passwords_and_errors,
        test_zip_corrupt_and_unsupported,
        test_homepage_collector,
        test_excel_merge_values_unchanged,
        test_file_inventory,
        test_rename_files,
        test_unreadable_workbook,
    ]
    for test in tests:
        print(f"\n[{test.__name__}]", flush=True)
        test()
    bad = [name for ok, name in RESULTS if not ok]
    print(f"\n{len(RESULTS) - len(bad)}/{len(RESULTS)} 통과")
    return 1 if bad else 0


if __name__ == "__main__":
    raise SystemExit(main())
